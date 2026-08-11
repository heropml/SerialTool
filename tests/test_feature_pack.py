# -*- coding: utf-8 -*-
"""Tests for dash_widgets / seq_report xlsx / Player drive_tx."""
from __future__ import print_function

import math
import os
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import dash_widgets
import rec_replay
import seq_report


class DashWidgetsTests(unittest.TestCase):
    def test_normalize_and_progress(self):
        self.assertEqual(dash_widgets.normalize_widget("GAUGE"), "gauge")
        self.assertEqual(dash_widgets.normalize_widget("nope"), "number")
        self.assertAlmostEqual(dash_widgets.progress_ratio(50, 0, 100), 0.5)
        self.assertEqual(dash_widgets.progress_ratio(-10, 0, 100), 0.0)
        self.assertEqual(dash_widgets.progress_ratio(200, 0, 100), 1.0)
        self.assertTrue(math.isfinite(dash_widgets.gauge_angle_deg(50, 0, 100)))
        # One-sided: do not invent a forced 0..100 opposite bound.
        self.assertAlmostEqual(dash_widgets.progress_ratio(50, None, 100), 0.5)
        self.assertAlmostEqual(dash_widgets.progress_ratio(20, 10, None), 1.0)

    def test_led_state(self):
        self.assertEqual(dash_widgets.led_state("", False), "ok")
        self.assertEqual(dash_widgets.led_state("warn", False), "warn")
        self.assertEqual(dash_widgets.led_state("", True), "alarm")
        self.assertEqual(dash_widgets.led_state("alarm", False), "alarm")

    def test_fmt_nan_inf(self):
        # Import helper from dashboard (requires Qt)
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from dashboard_dialog import _fmt
        self.assertEqual(_fmt(float("nan")), "--")
        self.assertEqual(_fmt(float("inf")), "--")
        self.assertEqual(_fmt(3.0), "3")


class SeqReportXlsxTests(unittest.TestCase):
    def test_report_fmt_xlsx(self):
        self.assertEqual(seq_report.report_fmt("a.xlsx"), ("xlsx", "a.xlsx"))
        self.assertEqual(
            seq_report.report_fmt("a", "Excel (*.xlsx)"),
            ("xlsx", "a.xlsx"))

    def test_build_xlsx_no_csv_safe_apostrophe(self):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest("openpyxl not installed")
        with tempfile.TemporaryDirectory() as td:
            path = os.path.join(td, "r.xlsx")
            seq_report.build_xlsx(
                "Title",
                [("k", "v")],
                "summary",
                ["A", "B"],
                [{"cells": ["=1+1", "ok"], "cls": "ok"}],
                path,
            )
            self.assertTrue(os.path.isfile(path))
            from openpyxl import load_workbook
            wb = load_workbook(path)
            ws = wb.active
            # Must NOT show a visible CSV-style leading apostrophe.
            found_formula_like = None
            for row in ws.iter_rows(min_row=1, max_col=2, values_only=False):
                for cell in row:
                    if cell.value == "=1+1":
                        found_formula_like = cell
                        break
                if found_formula_like is not None:
                    break
            self.assertIsNotNone(found_formula_like)
            self.assertFalse(str(found_formula_like.value).startswith("'"))
            self.assertEqual(found_formula_like.data_type, "s")


class PlayerDriveTxTests(unittest.TestCase):
    def test_drive_tx_filters_rx(self):
        got = []
        events = [
            (0.0, "rx", b"RX"),
            (0.1, "tx", b"TX1"),
            (0.2, "tx", b"TX2"),
        ]
        p = rec_replay.Player(
            events, None, mode="drive_tx", send_tx=got.append, speed=100.0)
        self.assertEqual(len(p), 2)
        p.start(0.0)
        p.tick(1.0)
        self.assertEqual(got, [b"TX1", b"TX2"])
        self.assertTrue(p.finished)
        self.assertEqual(p.send_fail_count, 0)

    def test_inject_mode_unchanged(self):
        got = []
        events = [(0.0, "rx", b"A"), (0.1, "tx", b"B")]
        p = rec_replay.Player(events, got.append, include_tx=False)
        self.assertEqual(len(p), 1)
        p.start(0.0)
        p.tick(1.0)
        self.assertEqual(got, [b"A"])

    def test_drive_tx_counts_failures_and_pauses(self):
        calls = []

        def boom(b):
            calls.append(b)
            return 0

        events = [(0.0, "tx", b"A"), (0.1, "tx", b"B"),
                  (0.2, "tx", b"C"), (0.3, "tx", b"D")]
        p = rec_replay.Player(
            events, None, mode="drive_tx", send_tx=boom, speed=100.0)
        p.start(0.0)
        p.tick(1.0)
        self.assertEqual(p.send_fail_count, 3)
        self.assertEqual(len(calls), 3)  # stop same tick after abort threshold
        self.assertTrue(p.paused)
        self.assertTrue(p.send_aborted)
        self.assertFalse(p.finished)

    def test_drive_tx_partial_write_counts_as_fail(self):
        def partial(b):
            return max(0, len(b) - 1)

        p = rec_replay.Player(
            [(0.0, "tx", b"ABCD")], None, mode="drive_tx",
            send_tx=partial, speed=100.0)
        p.start(0.0)
        p.tick(1.0)
        self.assertEqual(p.send_fail_count, 1)
        self.assertEqual(p.send_ok_count, 0)

    def test_drive_tx_abort_on_last_frame_finishes(self):
        def boom(_b):
            return 0

        events = [(0.0, "tx", b"A"), (0.1, "tx", b"B"), (0.2, "tx", b"C")]
        p = rec_replay.Player(
            events, None, mode="drive_tx", send_tx=boom, speed=100.0)
        p.start(0.0)
        p.tick(1.0)
        self.assertTrue(p.send_aborted)
        self.assertTrue(p.paused)
        self.assertTrue(p.finished)  # last frame reached; dialog must prefer finish cleanup


if __name__ == "__main__":
    unittest.main()
