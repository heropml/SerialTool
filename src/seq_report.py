# -*- coding: utf-8 -*-
"""Sequence report HTML/CSV renderers (Qt-free).

S-2 slice: formatters take already-localized strings so dialogs/CLI can share
the same output without PyQt.
"""
from __future__ import print_function

import csv
import html as _html
import io


def csv_safe(value):
    """Prefix formula-like CSV cells so Excel opens them as text."""
    if not isinstance(value, str):
        return value
    probe = value.lstrip()
    if value and (value[0] in "=+-@\t\r\n" or (probe and probe[0] in "=+-@")):
        return "'" + value
    return value


def report_fmt(path, sel=""):
    """Resolve export format + ensure extension. Returns (fmt, path)."""
    low = (path or "").lower()
    if low.endswith(".xlsx"):
        return "xlsx", path
    if low.endswith(".csv"):
        return "csv", path
    if low.endswith(".xml"):
        return "junit", path
    if low.endswith(".html") or low.endswith(".htm"):
        return "html", path
    sel_l = (sel or "").lower()
    if "xlsx" in sel_l or "excel" in sel_l:
        return "xlsx", path + ".xlsx"
    if "csv" in sel_l:
        return "csv", path + ".csv"
    if "xml" in sel_l or "junit" in sel_l:
        return "junit", path + ".xml"
    return "html", path + ".html"


def build_html(title, meta_rows, summary_line, passed, header, body):
    """Render a self-contained HTML report.

    body: iterable of {"cells": [...], "cls": "ok"|"fail"|""}
    """
    th = "".join("<th>%s</th>" % _html.escape(str(x)) for x in (header or []))
    trs = []
    for b in body or []:
        tds = "".join("<td>%s</td>" % _html.escape(str(x)) for x in b.get("cells", []))
        trs.append('<tr class="%s">%s</tr>' % (b.get("cls", ""), tds))
    meta_html = "".join(
        "<div class='meta'>%s: %s</div>" % (_html.escape(str(k)), _html.escape(str(v)))
        for k, v in (meta_rows or [])
    )
    summ = summary_line or ""
    return (
        "<!doctype html><html><head><meta charset='utf-8'><title>%(title)s</title><style>"
        "body{font-family:'Segoe UI','Microsoft YaHei',sans-serif;margin:24px;color:#222;}"
        "h1{font-size:20px;margin:0 0 6px;}"
        ".meta{color:#666;font-size:13px;margin-bottom:6px;}"
        ".verdict{display:inline-block;padding:3px 12px;border-radius:6px;color:#fff;"
        "font-weight:600;background:%(accent)s;}"
        "table{border-collapse:collapse;width:100%%;font-size:13px;margin-top:12px;}"
        "th,td{border:1px solid #ddd;padding:6px 8px;text-align:left;vertical-align:top;"
        "word-break:break-all;}"
        "th{background:#f4f5f7;font-weight:600;}"
        "tr.ok td{background:#ebfbee;}tr.fail td{background:#fff0f0;}"
        "td:first-child,th:first-child{text-align:center;width:36px;}"
        ".foot{color:#aaa;font-size:11px;margin-top:16px;}"
        "</style></head><body>"
        "<h1>%(title)s</h1>"
        "%(meta)s"
        "%(verdict_html)s"
        "<table><thead><tr>%(th)s</tr></thead><tbody>%(rows)s</tbody></table>"
        "<div class='foot'>CommTool - %(title)s</div></body></html>"
    ) % {
        "title": _html.escape(str(title or "")),
        "meta": meta_html,
        "verdict_html": (
            "<div><span class='verdict'>%s</span></div>" % _html.escape(summ)
        ) if summ else "",
        "accent": "#2f9e44" if passed else "#e03131",
        "th": th,
        "rows": "".join(trs),
    }


def build_csv(title, meta_rows, summary_line, header, body):
    """Render CSV text (Excel-safe cells)."""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([csv_safe(str(title or ""))])
    for k, v in (meta_rows or []):
        w.writerow([csv_safe(k), csv_safe(v)])
    if summary_line:
        w.writerow([csv_safe(summary_line)])
    w.writerow([])
    w.writerow([csv_safe(x) for x in (header or [])])
    for b in body or []:
        w.writerow([csv_safe(x) for x in b.get("cells", [])])
    return buf.getvalue()


def _xlsx_set(cell, value):
    """Write a cell for Excel without CSV apostrophe pollution.

    openpyxl does not evaluate formulas the way CSV+Excel does, so csv_safe()'s
    leading ``'`` would show up literally. Force string data_type for
    formula-like text so Excel still won't auto-run them.
    """
    if value is None:
        cell.value = ""
        return
    if isinstance(value, bool):
        cell.value = value
        return
    if isinstance(value, (int, float)):
        cell.value = value
        return
    text = str(value)
    cell.value = text
    if text[:1] in "=+-@\t\r\n":
        cell.data_type = "s"


def build_xlsx(title, meta_rows, summary_line, header, body, path):
    """Write an .xlsx workbook mirroring build_csv layout. Requires openpyxl."""
    try:
        from openpyxl import Workbook
    except ImportError as e:
        raise RuntimeError("openpyxl is required for Excel export") from e

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    row_i = 1
    _xlsx_set(ws.cell(row=row_i, column=1), str(title or ""))
    row_i += 1
    for k, v in (meta_rows or []):
        _xlsx_set(ws.cell(row=row_i, column=1), k)
        _xlsx_set(ws.cell(row=row_i, column=2), v)
        row_i += 1
    if summary_line:
        _xlsx_set(ws.cell(row=row_i, column=1), summary_line)
        row_i += 1
    row_i += 1
    for col, h in enumerate(header or [], 1):
        _xlsx_set(ws.cell(row=row_i, column=col), h)
    row_i += 1
    for b in body or []:
        for col, cell in enumerate(b.get("cells", []), 1):
            _xlsx_set(ws.cell(row=row_i, column=col), cell)
        row_i += 1
    wb.save(path)
    return path
